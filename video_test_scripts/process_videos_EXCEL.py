import os
import cv2
import time
import torch
import openpyxl
import numpy as np
from tqdm import tqdm
from ultralytics import YOLO

NEW_EXCEL_FILE_NAME = "VideoResultsNEW.xlsx"

FILE_NAME_LIGHT_CLEAN_25 = "op25_light_clean.mp4"
FILE_NAME_LIGHT_CLEAN_50 = "op50_light_clean.mp4"
FILE_NAME_LIGHT_CLEAN_75 = "op75_light_clean.mp4"
FILE_NAME_LIGHT_CLEAN_100 = "op100_light_clean.mp4"

FILE_NAME_DARK_CLEAN_25 = "op25_dark_clean.mp4"
FILE_NAME_DARK_CLEAN_50 = "op50_dark_clean.mp4"
FILE_NAME_DARK_CLEAN_75 = "op75_dark_clean.mp4"
FILE_NAME_DARK_CLEAN_100 = "op100_dark_clean.mp4"

FILE_NAME_LIGHT_NOISY_25 = "op25_light_noisy.mp4"
FILE_NAME_LIGHT_NOISY_50 = "op50_light_noisy.mp4"
FILE_NAME_LIGHT_NOISY_75 = "op75_light_noisy.mp4"
FILE_NAME_LIGHT_NOISY_100 = "op100_light_noisy.mp4"

FILE_NAME_DARK_NOISY_25 = "op25_dark_noisy.mp4"
FILE_NAME_DARK_NOISY_50 = "op50_dark_noisy.mp4"
FILE_NAME_DARK_NOISY_75 = "op75_dark_noisy.mp4"
FILE_NAME_DARK_NOISY_100 = "op100_dark_noisy.mp4"

VIDEO_FILES_1 = [FILE_NAME_LIGHT_CLEAN_25, FILE_NAME_LIGHT_CLEAN_50, FILE_NAME_LIGHT_CLEAN_75, FILE_NAME_LIGHT_CLEAN_100]
VIDEO_FILES_2 = [FILE_NAME_DARK_CLEAN_25, FILE_NAME_DARK_CLEAN_50, FILE_NAME_DARK_CLEAN_75, FILE_NAME_DARK_CLEAN_100]
VIDEO_FILES_3 = [FILE_NAME_LIGHT_NOISY_25, FILE_NAME_LIGHT_NOISY_50, FILE_NAME_LIGHT_NOISY_75, FILE_NAME_LIGHT_NOISY_100]
VIDEO_FILES_4 = [FILE_NAME_DARK_NOISY_25, FILE_NAME_DARK_NOISY_50, FILE_NAME_DARK_NOISY_75, FILE_NAME_DARK_NOISY_100]


def create_spreadsheet():
    wb = openpyxl.load_workbook('video_test_scripts/VideoResultsTEMPLATE.xlsx')
    wb.save(NEW_EXCEL_FILE_NAME)
    wb.close()


def process_with_YOLO():
    model = YOLO("yolo11n.pt")

    for video_set_num in range(1, 5):
        if video_set_num == 1:
            tqdm_desc = "[Part 1/8] YOLO Processing  "
            video_files_set = VIDEO_FILES_1
            row_start_index = 4
            col_start_index = 3
        elif video_set_num == 2:
            tqdm_desc = "[Part 2/8] YOLO Processing  "
            video_files_set = VIDEO_FILES_2
            row_start_index = 10
            col_start_index = 3
        elif video_set_num == 3:
            tqdm_desc = "[Part 3/8] YOLO Processing  "
            video_files_set = VIDEO_FILES_3
            row_start_index = 4
            col_start_index = 8
        elif video_set_num == 4:
            tqdm_desc = "[Part 4/8] YOLO Processing  "
            video_files_set = VIDEO_FILES_4
            row_start_index = 10
            col_start_index = 8
        else:
            print(f"Error: Invalid video files set number: {video_set_num}")
            return

        # loop through video files in set
        for i in tqdm(range(len(video_files_set)), desc=tqdm_desc):
            fname = video_files_set[i]

            # read video file
            cap = cv2.VideoCapture(os.path.join("videos", fname))
            if not cap.isOpened():
                print(f"Error: Could not open video file: {fname}")
                return
            
            # total frames in video
            total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
            # print(f"Total frames in video: {total_frames}")

            frame_counter = 0
            detection_counter = 0
            
            # loop through video frames
            start_time = time.time()
            while cap.isOpened():
                ret, frame = cap.read()
                if not ret:
                    print("Error: Can't receive frame")
                    break

                frame_counter += 1
                
                # get YOLO detection results
                result = model.predict(frame, conf=0.5, classes=[67], verbose=False)[0].boxes.xywh.cpu().numpy()
                if len(result) > 0:
                    for box in result:
                        x, y, w, h = box
                        cv2.rectangle(frame, (int(x-(w/2)), int(y-(h/2))), (int(x+(w/2)), int(y+(h/2))), (0, 255, 0), 2)
                        detection_counter += 1

                if frame_counter >= total_frames:
                    break

            # release resources
            cap.release()
            elapsed_time = time.time() - start_time
            results = [
                np.round(frame_counter / elapsed_time, 2), # Avg FPS
                np.round(detection_counter / elapsed_time, 2), # Avg DPS
                frame_counter, # Total num frames
                detection_counter, # Total num detections
                np.round(elapsed_time, 2) # Total time
            ]

            # write results to excel file
            wb = openpyxl.load_workbook(NEW_EXCEL_FILE_NAME)
            ws = wb.active
            for j in range(len(results)):
                ws.cell(row=row_start_index + i, column=col_start_index + j, value=results[j])
            wb.save(NEW_EXCEL_FILE_NAME)
            wb.close()


def process_with_aruco():
    aruco_dict = cv2.aruco.getPredefinedDictionary(cv2.aruco.DICT_6X6_250)
    aruco_params =  cv2.aruco.DetectorParameters()
    aruco_detector = cv2.aruco.ArucoDetector(aruco_dict, aruco_params)

    for video_set_num in range(1, 5):
        if video_set_num == 1:
            tqdm_desc = "[Part 5/8] ArUco Processing "
            video_files_set = VIDEO_FILES_1
            row_start_index = 16
            col_start_index = 3
        elif video_set_num == 2:
            tqdm_desc = "[Part 6/8] ArUco Processing "
            video_files_set = VIDEO_FILES_2
            row_start_index = 22
            col_start_index = 3
        elif video_set_num == 3:
            tqdm_desc = "[Part 7/8] ArUco Processing "
            video_files_set = VIDEO_FILES_3
            row_start_index = 16
            col_start_index = 8
        elif video_set_num == 4:
            tqdm_desc = "[Part 8/8] ArUco Processing "
            video_files_set = VIDEO_FILES_4
            row_start_index = 22
            col_start_index = 8
        else:
            print(f"Error: Invalid video files set number: {video_set_num}")
            return

        # loop through video files in set
        for i in tqdm(range(len(video_files_set)), desc=tqdm_desc):
            fname = video_files_set[i]

            # read video file
            cap = cv2.VideoCapture(os.path.join("videos", fname))
            if not cap.isOpened():
                print(f"Error: Could not open video file: {fname}")
                return
            
            total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
            frame_counter = 0
            detection_counter = 0
            
            # loop through video frames
            start_time = time.time()
            while cap.isOpened():
                ret, frame = cap.read()
                if not ret:
                    print("Error: Can't receive frame")
                    break

                frame_counter += 1
                
                # get aruco detection results
                _, ids, _ = aruco_detector.detectMarkers(frame)
                if ids is not None:
                    detection_counter += len(ids)

                if frame_counter >= total_frames:
                    break

            # release resources 
            cap.release()
            elapsed_time = time.time() - start_time
            results = [
                np.round(frame_counter / elapsed_time, 2), # Avg FPS
                np.round(detection_counter / elapsed_time, 2), # Avg DPS
                frame_counter, # Total num frames
                detection_counter, # Total num detections
                np.round(elapsed_time, 2) # Total time
            ]

            # write results to excel file
            wb = openpyxl.load_workbook(NEW_EXCEL_FILE_NAME)
            ws = wb.active
            for j in range(len(results)):
                ws.cell(row=row_start_index + i, column=col_start_index + j, value=results[j])
            wb.save(NEW_EXCEL_FILE_NAME)
            wb.close()


if __name__ == '__main__':
    try:
        create_spreadsheet()
        process_with_YOLO()
        process_with_aruco()
    except KeyboardInterrupt:
        cv2.destroyAllWindows()
